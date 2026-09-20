import io
from datetime import datetime, timezone

# ---------------------------------------------------------------------------
# Contenu initial (source de vérité data-driven — ajouter un produit ne demande
# pas de toucher au code des pages).
# ---------------------------------------------------------------------------

SETTINGS = {
    "id": "site",
    "brand_name": "La Maison d'Auben",
    "tagline": "L'aubaine pour mieux s'organiser.",
    "slogans": [
        "L'aubaine pour mieux s'organiser.",
        "Des outils qui font la différence.",
        "De bonnes idées. De bons outils. Une belle aubaine.",
    ],
    "hero_intro": "Des outils numériques simples, beaux et intelligents pour mieux organiser votre vie, vos projets et vos envies. Le premier d'une longue série.",
    "newsletter_title": "Recevez nos nouveaux outils et ressources gratuitement.",
    "newsletter_text": "Pas de spam, juste nos meilleures idées pour mieux vous organiser, et nos nouveautés en avant-première.",
    "currency": "eur",
}

CATEGORIES = [
    {"slug": "budget-finance", "name": "Budget & Finance", "icon": "Wallet", "order": 1,
     "description": "Reprenez le contrôle de votre argent, simplement."},
    {"slug": "travel", "name": "Travel", "icon": "Plane", "order": 2,
     "description": "Préparez vos voyages et roadtrips l'esprit léger."},
    {"slug": "maison", "name": "Maison & Vie quotidienne", "icon": "Home", "order": 3,
     "description": "Organisez votre foyer et votre quotidien."},
    {"slug": "automobile", "name": "Automobile", "icon": "Car", "order": 4,
     "description": "Suivez le vrai coût de votre véhicule."},
    {"slug": "investissement", "name": "Investissement", "icon": "TrendingUp", "order": 5,
     "description": "Suivez et comprenez vos placements."},
    {"slug": "couple", "name": "Couple", "icon": "Heart", "order": 6,
     "description": "Gérez vos projets et votre budget à deux."},
    {"slug": "lifestyle", "name": "Lifestyle", "icon": "Backpack", "order": 7,
     "description": "Des outils pour vos projets de vie."},
    {"slug": "loisirs", "name": "Loisirs", "icon": "Dices", "order": 8,
     "description": "Collections, randonnée, JDR et plus encore."},
]

GALLERY_SLOTS = [
    {"key": "dashboard", "caption": "Tableau de bord", "placeholder": True},
    {"key": "revenus", "caption": "Revenus", "placeholder": True},
    {"key": "depenses", "caption": "Dépenses", "placeholder": True},
    {"key": "historique", "caption": "Historique des opérations", "placeholder": True},
    {"key": "analyse", "caption": "Analyse du budget", "placeholder": True},
    {"key": "graphiques", "caption": "Graphiques", "placeholder": True},
]

PRODUCTS = [
    {
        "slug": "budget-mensuel",
        "name": "Budget mensuel",
        "category_slug": "budget-finance",
        "lookup_key": "budget_mensuel",
        "price": 7.99,
        "currency": "eur",
        "badge": "Nouveau",
        "status": "available",
        "short_description": "Le tableau de bord simple pour reprendre le contrôle de son budget.",
        "description": "Un outil simple et complet pour suivre ses revenus, ses dépenses, son épargne et ses objectifs, et mieux comprendre son budget au quotidien.",
        "audience": "Que vous soyez étudiant, jeune actif, en couple ou en famille, Budget mensuel vous aide à voir clair dans vos finances sans être un expert d'Excel. Aucune connaissance technique requise.",
        "features": [
            {"icon": "Wallet", "title": "Revenus", "text": "Enregistrez toutes vos rentrées d'argent en un coup d'œil."},
            {"icon": "Receipt", "title": "Dépenses", "text": "Dépenses fixes et variables, catégorisées automatiquement."},
            {"icon": "PiggyBank", "title": "Épargne", "text": "Suivez votre épargne mois après mois."},
            {"icon": "Target", "title": "Objectifs", "text": "Fixez des objectifs et mesurez votre progression."},
            {"icon": "Calendar", "title": "Suivi mensuel", "text": "Une vue claire pour chaque mois de l'année."},
            {"icon": "BarChart3", "title": "Analyse annuelle", "text": "Graphiques et indicateurs pour comprendre vos tendances."},
        ],
        "contents": [
            "Un fichier Excel prêt à l'emploi (compatible Google Sheets)",
            "Un tableau de bord automatisé avec graphiques",
            "Des feuilles Revenus, Dépenses, Épargne, Objectifs",
            "Un suivi mensuel et annuel",
            "Un guide de démarrage rapide",
        ],
        "compatibility": [
            "Microsoft Excel (Windows & Mac)",
            "Google Sheets",
            "LibreOffice Calc",
        ],
        "faq": [
            {"q": "Ai-je besoin de connaître Excel ?", "a": "Non. Le fichier est pensé pour les débutants : vous remplissez, tout se calcule automatiquement."},
            {"q": "Le fichier fonctionne-t-il sur Mac ?", "a": "Oui, il fonctionne sur Excel Windows et Mac, ainsi que sur Google Sheets."},
            {"q": "Puis-je le personnaliser ?", "a": "Bien sûr. Catégories, couleurs et intitulés sont entièrement modifiables."},
            {"q": "Comment je reçois le fichier ?", "a": "Immédiatement après le paiement, via un lien de téléchargement sécurisé et par e-mail."},
        ],
        "seo": {
            "title": "Budget mensuel — Template Excel pour gérer son budget | La Maison d'Auben",
            "description": "Un template Excel simple et complet pour suivre vos revenus, dépenses, épargne et objectifs. Reprenez le contrôle de votre budget dès 7,99 €.",
        },
        "gallery": GALLERY_SLOTS,
        "main_image_placeholder": True,
    },
]

# Produits « Bientôt disponible » (ambition de la marque, sans tromper l'utilisateur)
COMING_SOON = [
    {"slug": "travel-manager", "name": "Travel Manager", "category_slug": "travel", "status": "coming_soon",
     "short_description": "Planifiez chaque voyage, du budget à l'itinéraire.", "price": None},
    {"slug": "car-cost", "name": "Car Cost", "category_slug": "automobile", "status": "coming_soon",
     "short_description": "Le vrai coût de votre voiture, enfin clair.", "price": None},
    {"slug": "first-apartment", "name": "First Apartment", "category_slug": "maison", "status": "coming_soon",
     "short_description": "Préparez votre premier appartement sereinement.", "price": None},
    {"slug": "budget-couple", "name": "Budget Couple", "category_slug": "couple", "status": "coming_soon",
     "short_description": "Gérez votre budget à deux, en toute transparence.", "price": None},
]

ARTICLES = [
    {
        "slug": "comment-faire-son-budget-mensuel",
        "title": "Comment faire son budget mensuel ?",
        "category": "Budget & Finance",
        "excerpt": "La méthode simple pour poser son budget en 30 minutes, même quand on part de zéro.",
        "read_time": 6,
        "cover_image": "https://images.unsplash.com/photo-1564510714747-69c3bc1fab41?crop=entropy&cs=srgb&fm=jpg&q=85&w=1200",
        "related_product_slugs": ["budget-mensuel"],
        "published_at": "2026-01-15T09:00:00+00:00",
        "content": [
            {"type": "p", "text": "Faire son budget n'a rien de compliqué. L'objectif n'est pas de tout contrôler au centime près, mais de voir clair : combien entre, combien sort, et ce qu'il reste."},
            {"type": "h2", "text": "1. Faites la liste de vos revenus"},
            {"type": "p", "text": "Commencez par le plus simple : tout l'argent qui rentre chaque mois. Salaire, aides, revenus complémentaires. Notez le montant net, celui que vous recevez réellement."},
            {"type": "h2", "text": "2. Séparez dépenses fixes et variables"},
            {"type": "p", "text": "Les dépenses fixes (loyer, abonnements, assurances) tombent chaque mois. Les variables (courses, sorties, imprévus) bougent. Les distinguer change tout."},
            {"type": "h2", "text": "3. Fixez un objectif d'épargne"},
            {"type": "p", "text": "Même petit, un objectif régulier fait la différence. Payez-vous en premier : mettez de côté dès le début du mois, pas avec ce qu'il reste à la fin."},
        ],
        "cta_text": "Vous souhaitez aller plus loin ? Découvrez Budget mensuel.",
    },
    {
        "slug": "combien-coute-reellement-une-voiture",
        "title": "Combien coûte réellement une voiture ?",
        "category": "Automobile",
        "excerpt": "Au-delà du prix d'achat, voici tous les coûts qu'on oublie trop souvent.",
        "read_time": 7,
        "cover_image": "https://images.unsplash.com/photo-1449965408869-eaa3f722e40d?crop=entropy&cs=srgb&fm=jpg&q=85&w=1200",
        "related_product_slugs": [],
        "published_at": "2026-01-10T09:00:00+00:00",
        "content": [
            {"type": "p", "text": "Le prix affiché n'est que la partie visible de l'iceberg. Une voiture coûte bien plus que son ticket d'achat."},
            {"type": "h2", "text": "Les coûts cachés"},
            {"type": "p", "text": "Assurance, carburant, entretien, pneus, décote, stationnement... Additionnés sur une année, ces postes dépassent souvent la mensualité de crédit."},
        ],
        "cta_text": "Car Cost arrive bientôt dans la Maison.",
    },
    {
        "slug": "comment-commencer-a-epargner",
        "title": "Comment commencer à épargner ?",
        "category": "Budget & Finance",
        "excerpt": "Trois habitudes simples pour se constituer une épargne, sans se priver.",
        "read_time": 5,
        "cover_image": "https://images.pexels.com/photos/8251157/pexels-photo-8251157.jpeg?auto=compress&cs=tinysrgb&w=1200",
        "related_product_slugs": ["budget-mensuel"],
        "published_at": "2026-01-05T09:00:00+00:00",
        "content": [
            {"type": "p", "text": "Épargner, ce n'est pas gagner plus. C'est surtout organiser ce que l'on a déjà."},
            {"type": "h2", "text": "Automatisez"},
            {"type": "p", "text": "Programmez un virement automatique le jour de la paie. Ce que vous ne voyez pas, vous ne le dépensez pas."},
        ],
        "cta_text": "Vous souhaitez aller plus loin ? Découvrez Budget mensuel.",
    },
]

GLOBAL_FAQ = [
    {"question": "Comment vais-je recevoir mon produit ?", "answer": "Dès votre paiement validé, vous accédez immédiatement à un lien de téléchargement sécurisé, et vous recevez également un e-mail de confirmation avec ce lien.", "order": 1},
    {"question": "Les paiements sont-ils sécurisés ?", "answer": "Oui. Les paiements sont traités par Stripe, un acteur de référence mondial. Nous ne stockons aucune donnée bancaire.", "order": 2},
    {"question": "Ai-je besoin d'un compte pour acheter ?", "answer": "Non. Vous pouvez acheter et télécharger votre produit sans créer de compte.", "order": 3},
    {"question": "Puis-je être remboursé ?", "answer": "Nos produits étant numériques et téléchargeables immédiatement, les conditions de remboursement sont précisées dans nos CGV. Contactez-nous en cas de problème, nous cherchons toujours une solution.", "order": 4},
    {"question": "Sur quels logiciels fonctionnent les fichiers ?", "answer": "Nos templates fonctionnent avec Microsoft Excel, Google Sheets et LibreOffice Calc. La compatibilité précise est indiquée sur chaque fiche produit.", "order": 5},
]


def build_placeholder_xlsx() -> bytes:
    """Fichier Excel PLACEHOLDER, clairement identifié comme tel, remplaçable
    par le vrai fichier via l'administration."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "Budget mensuel"
    green = PatternFill(start_color="1E3A2B", end_color="1E3A2B", fill_type="solid")
    ws["B2"] = "La Maison d'Auben — Budget mensuel"
    ws["B2"].font = Font(bold=True, size=16, color="1E3A2B")
    ws["B4"] = "FICHIER PLACEHOLDER"
    ws["B4"].font = Font(bold=True, color="FFFFFF")
    ws["B4"].fill = green
    ws["B6"] = "Ceci est un fichier de démonstration."
    ws["B7"] = "Remplacez-le par le vrai fichier Excel depuis l'administration."
    for i, h in enumerate(["Mois", "Revenus", "Dépenses fixes", "Dépenses variables", "Épargne"]):
        c = ws.cell(row=10, column=2 + i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = green
        c.alignment = Alignment(horizontal="center")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
